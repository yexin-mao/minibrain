from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import Workbook

from minibrain import gateway
from minibrain.contracts import ModuleError
from minibrain.db import table_db
from minibrain.modules.table_rag import core
from minibrain.modules.table_rag.xlsx import extract_sheet, inspect_workbook, is_xlsx


def _workbook_bytes(build=None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    if build:
        build(workbook)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _sales_workbook() -> bytes:
    def build(workbook):
        sales = workbook["销售"]
        sales.append(["地区", "金额", "日期", "有效", "编号", "地区"])
        sales.append(["华东", 1200, dt.date(2026, 8, 1), True, "00123", "一部"])
        sales.append(["华北", 800, dt.date(2026, 8, 2), False, "00456", "二部"])
        inventory = workbook.create_sheet("库存")
        inventory.append(["产品", "数量"])
        inventory.append(["A", 3])
        hidden = workbook.create_sheet("内部计算")
        hidden.append(["不应导入", "值"])
        hidden.append(["secret", 1])
        hidden.sheet_state = "hidden"

    return _workbook_bytes(build)


def test_xlsx_is_detected_by_package_structure_not_extension():
    raw = _sales_workbook()
    assert is_xlsx(raw)
    assert core.detect_upload_kind(object(), "renamed.bin", raw) == "xlsx"


def test_only_visible_sheets_are_registered():
    info = inspect_workbook(_sales_workbook(), "sales.xlsx")
    assert info.visible_sheets == ["销售", "库存"]


def test_sheet_extraction_preserves_typed_values_as_importable_text():
    rows = extract_sheet(_sales_workbook(), "sales.xlsx", "销售")
    assert rows[0] == ["地区", "金额", "日期", "有效", "编号", "地区"]
    assert rows[1] == ["华东", "1200", "2026-08-01", "TRUE", "00123", "一部"]


def test_uncalculated_formula_is_rejected_instead_of_becoming_blank():
    def build(workbook):
        sheet = workbook["销售"]
        sheet.append(["金额", "含税"])
        sheet.append([100, "=A2*1.13"])

    with pytest.raises(ModuleError) as caught:
        extract_sheet(_workbook_bytes(build), "formula.xlsx", "销售")
    assert caught.value.code == "xlsx_uncalculated_formula"
    assert "重新计算并保存" in caught.value.message


def test_xlsx_upload_creates_one_ready_dataset_per_visible_sheet(alice):
    ids = gateway.call(
        "table-rag", "upload_xlsx", alice, None, "sales.xlsx", _sales_workbook())
    assert len(ids) == 2
    for dataset_id in ids:
        gateway.process("table-rag", dataset_id)

    datasets = [item for item in gateway.call("table-rag", "list_datasets", alice)
                if str(item["id"]) in ids]
    assert {item["sheet_name"] for item in datasets} == {"销售", "库存"}
    assert all(item["status"] == "ready" for item in datasets)
    assert all(item["parsed_as"] == "xlsx" for item in datasets)
    assert {item["row_count"] for item in datasets} == {1, 2}
    assert len({str(item["workbook_id"]) for item in datasets}) == 1

    sales = next(item for item in datasets if item["sheet_name"] == "销售")
    types = {column["name"]: column["type"] for column in sales["columns"]}
    assert types == {
        "地区": "text", "金额": "numeric", "日期": "date",
        "有效": "boolean", "编号": "text", "地区_2": "text",
    }
    result = gateway.call(
        "table-rag", "run_query", alice,
        f'SELECT sum("金额") AS total FROM {sales["table_name"]}',
    )
    assert float(result["rows"][0]["total"]) == 2000


def test_deleting_last_sheet_removes_shared_workbook(alice):
    ids = gateway.call(
        "table-rag", "upload_xlsx", alice, None, "delete.xlsx", _sales_workbook())
    rows = [item for item in gateway.call("table-rag", "list_datasets", alice)
            if str(item["id"]) in ids]
    workbook_id = str(rows[0]["workbook_id"])

    gateway.call("table-rag", "delete_dataset", alice, ids[0])
    with table_db() as cur:
        cur.execute("SELECT 1 FROM workbooks WHERE id = %s", (workbook_id,))
        assert cur.fetchone() is not None

    gateway.call("table-rag", "delete_dataset", alice, ids[1])
    with table_db() as cur:
        cur.execute("SELECT 1 FROM workbooks WHERE id = %s", (workbook_id,))
        assert cur.fetchone() is None
