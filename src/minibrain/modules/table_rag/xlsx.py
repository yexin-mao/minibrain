"""XLSX 内容识别、压缩包护栏和单 Sheet 提取。"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal

from ...contracts import ModuleError

MAX_ARCHIVE_ENTRIES = 5000
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_VISIBLE_SHEETS = 20
MAX_ROWS_PER_SHEET = 100_000
MAX_COLUMNS_PER_SHEET = 256
MAX_CELLS_PER_SHEET = 1_000_000


@dataclass(frozen=True)
class WorkbookInfo:
    visible_sheets: list[str]


def _archive_names(raw: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            return set(archive.namelist())
    except (zipfile.BadZipFile, OSError):
        return set()


def is_xlsx(raw: bytes) -> bool:
    """按 OOXML 包结构识别，不相信扩展名。"""
    names = _archive_names(raw)
    return "xl/workbook.xml" in names and "[Content_Types].xml" in names


def _validate_archive(raw: bytes, filename: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            infos = archive.infolist()
            names = {item.filename for item in infos}
    except (zipfile.BadZipFile, OSError) as exc:
        raise ModuleError(f"{filename}：XLSX 压缩包损坏", code="bad_xlsx") from exc
    if "xl/workbook.xml" not in names:
        raise ModuleError(f"{filename} 不是有效 XLSX", code="bad_xlsx")
    if "xl/vbaProject.bin" in names:
        raise ModuleError("不支持含宏的 Excel 工作簿，请另存为 .xlsx", code="xlsx_macro")
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        raise ModuleError("XLSX 内部文件数量异常", code="xlsx_archive_limit")
    if sum(item.file_size for item in infos) > MAX_UNCOMPRESSED_BYTES:
        raise ModuleError("XLSX 解压后超过 200 MB 安全上限", code="xlsx_archive_limit")


def _load(raw: bytes, *, data_only: bool):
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        return load_workbook(
            io.BytesIO(raw), read_only=True, data_only=data_only,
            keep_links=False, keep_vba=False,
        )
    except Exception as exc:  # noqa: BLE001
        raise ModuleError(
            f"XLSX 解析失败：{type(exc).__name__}", code="bad_xlsx") from exc


def inspect_workbook(raw: bytes, filename: str) -> WorkbookInfo:
    """上传期只读取工作簿目录，不遍历单元格。"""
    _validate_archive(raw, filename)
    workbook = _load(raw, data_only=False)
    try:
        visible = [sheet.title for sheet in workbook.worksheets
                   if sheet.sheet_state == "visible"]
    finally:
        workbook.close()
    if not visible:
        raise ModuleError("XLSX 没有可见 Sheet", code="empty_xlsx")
    if len(visible) > MAX_VISIBLE_SHEETS:
        raise ModuleError(
            f"XLSX 可见 Sheet 超过 {MAX_VISIBLE_SHEETS} 个上限",
            code="xlsx_sheet_limit")
    return WorkbookInfo(visible_sheets=visible)


def _has_time_format(number_format: str) -> bool:
    """Excel 的 m 同时表示月/分钟；h 或 s 才能无歧义说明包含时间。"""
    cleaned = number_format.lower().replace('"', "").replace("\\", "")
    return "h" in cleaned or "s" in cleaned


def _cell_text(value: object, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if not _has_time_format(number_format):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    return str(value).strip()


def extract_sheet(raw: bytes, filename: str, sheet_name: str) -> list[list[str]]:
    """提取一个 Sheet；公式使用 Excel 已保存的缓存值，禁止静默吞掉未计算公式。"""
    _validate_archive(raw, filename)
    formulas = _load(raw, data_only=False)
    values = _load(raw, data_only=True)
    try:
        if sheet_name not in formulas.sheetnames or sheet_name not in values.sheetnames:
            raise ModuleError(f"Sheet 不存在：{sheet_name}", code="xlsx_sheet_missing")
        formula_sheet = formulas[sheet_name]
        value_sheet = values[sheet_name]
        if (formula_sheet.max_row > MAX_ROWS_PER_SHEET
                or formula_sheet.max_column > MAX_COLUMNS_PER_SHEET
                or formula_sheet.max_row * formula_sheet.max_column > MAX_CELLS_PER_SHEET):
            raise ModuleError(
                f"Sheet {sheet_name} 超过 {MAX_ROWS_PER_SHEET} 行、"
                f"{MAX_COLUMNS_PER_SHEET} 列或 {MAX_CELLS_PER_SHEET} 单元格上限",
                code="xlsx_sheet_size_limit")

        rows: list[list[str]] = []
        value_rows = value_sheet.iter_rows()
        for row_number, formula_row in enumerate(formula_sheet.iter_rows(), start=1):
            cached_row = next(value_rows)
            converted = []
            for formula_cell, cached_cell in zip(formula_row, cached_row):
                if formula_cell.data_type == "f" and cached_cell.value is None:
                    raise ModuleError(
                        f"Sheet {sheet_name} 的 {formula_cell.coordinate} 是未计算公式；"
                        "请用 Excel/LibreOffice 重新计算并保存后上传",
                        code="xlsx_uncalculated_formula")
                converted.append(_cell_text(cached_cell.value, cached_cell.number_format))
            while converted and converted[-1] == "":
                converted.pop()
            if any(converted):
                rows.append(converted)
            if row_number > MAX_ROWS_PER_SHEET:
                raise ModuleError("XLSX 行数超过安全上限", code="xlsx_sheet_size_limit")
        return rows
    finally:
        formulas.close()
        values.close()
